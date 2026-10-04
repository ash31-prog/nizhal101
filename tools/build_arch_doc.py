"""Builds docs/NIZHAL_Architecture.docx from the same content as the app."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

EBONY = RGBColor(0x1C,0x2E,0x40); TEAL = RGBColor(0x00,0x80,0x83)
RUST  = RGBColor(0xC9,0x51,0x0D); INK2 = RGBColor(0x4A,0x60,0x71)
INK3  = RGBColor(0x7C,0x8F,0x9D)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.85)
    s.top_margin = s.bottom_margin = Inches(0.8)
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

def H(text, size=20, color=EBONY, space=10, bold=True):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space); p.paragraph_format.space_before = Pt(space+4)
    return p

def P(text, size=10.5, color=None, italic=False, space=6):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.size = Pt(size); r.italic = italic
    if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space)
    return p

def bullets(items):
    for i in items:
        p = doc.add_paragraph(style='List Bullet'); r = p.add_run(i)
        r.font.size = Pt(10.5); p.paragraph_format.space_after = Pt(2)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), hexcolor); tcPr.append(sh)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); shade(c, '1C2E40')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(9.5)
    if widths:
        for r_ in t.rows:
            for i, w in enumerate(widths): r_.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t

def image(path, caption, width=6.6):
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption); r.font.size = Pt(8.5); r.italic = True; r.font.color.rgb = INK3
    c.paragraph_format.space_after = Pt(14)

def code(text):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.name = 'Consolas'; r.font.size = Pt(8.5); r.font.color.rgb = INK2
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.15)

# ---------------- title ----------------
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('NIZHAL'); r.bold = True; r.font.size = Pt(40); r.font.color.rgb = EBONY
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run('System Architecture & Deployment Specification')
r.font.size = Pt(14); r.font.color.rgb = TEAL
s2 = doc.add_paragraph(); s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s2.add_run('PS 08 · Safer Urban Mobility for Women & Vulnerable Citizens')
r.font.size = Pt(10.5); r.font.color.rgb = INK3
q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = q.add_run('“Every safety app asks you to press a button. We built one you never touch.”')
r.italic = True; r.font.size = Pt(12); r.font.color.rgb = EBONY
q.paragraph_format.space_before = Pt(18); q.paragraph_format.space_after = Pt(22)

H('1 · Executive summary', 16)
P('NIZHAL is a client-heavy, server-minimal system. A street-risk map is computed offline from '
  'open data and shipped inside the application. On the phone, five sensing layers produce twelve '
  'features; a fusion function and a trained autoencoder collapse those into a single score from 0 to 1, '
  'which drives a five-state machine. Only the terminal state transmits, and what it transmits is an '
  'AES-256-GCM encrypted blob that our own server cannot read.')
P('The architecture is itself the product argument: we are not asking to be trusted, we have made '
  'ourselves incapable of betraying that trust.', italic=True, color=TEAL)

table(['Metric', 'Value', 'Meaning'],
 [['Time to detection', 'under 90 seconds', 'with zero user action'],
  ['Users needed for the map', '0', 'computed, not crowdsourced'],
  ['Audio bytes transmitted', '0', 'analysis is on-device, buffer overwritten'],
  ['Offline capability', '100% of core alerting', 'service worker + SMS fallback'],
  ['External AI services', 'none', 'no API keys exist in this project']],
 [1.7, 1.6, 3.3])

doc.add_page_break()
H('2 · Layered architecture', 16)
P('Five layers. Solid borders are running in the prototype you can open today. '
  'Dashed borders are on the roadmap and are labelled as such everywhere, including in the product UI.')
image('docs/img/layers.png', 'Figure 1 — The five layers, with shipped vs roadmap status per component.', 6.6)

doc.add_page_break()
H('3 · Layer detail', 16)

H('A · City Risk Pipeline', 13, TEAL, 6)
P('Input is an Overpass extract of Thiruvanmiyur and Adyar, Chennai: 2,970 street segments, '
  '2,322 junctions, 212.5 km of road and 213 mapped amenities.')
code('r(segment, hour) = f(lighting, frontage, footpath, enclosure,\n'
     '                     topological isolation, segment length)\n'
     '                   × night and late-night multipliers')
P('Provenance is tracked per segment and surfaced in the dashboard. The OSM lit tag exists on 209 '
  'segments; frontage is measured from real points of interest on 835; the remainder uses a '
  'road-class prior. We display which is which rather than implying the whole map is measured.')
P('Output is a 477 KB bundle compiled into the application, which is why routing and risk scoring '
  'continue to work with the network switched off.')

H('B · On-Device Engine', 13, TEAL, 6)
bullets([
 'Motion — RMS, jerk and spectral entropy from a Goertzel bank over 0.5–12 Hz. Walking is periodic and scores low entropy; a struggle is broadband and aperiodic.',
 'Audio (opt-in) — fundamental frequency, share of energy in 1–4 kHz, and roughness (amplitude modulation). Two-second RAM ring buffer, continuously overwritten, never stored or transmitted.',
 'Anomaly — a 9-5-3-5-9 dense autoencoder trained on normal behaviour only: 24,000 samples of walking, vehicle rides, pocket jostle and stationary use. A calm walk scores 0.01; being pulled into a vehicle scores 1.00. No assault data was used, or needed.',
 'Fusion — twelve features become weighted evidence, saturated into 0–1, multiplied by a context term derived from street risk, hour and isolation.',
 'States — OFF, ARMED, HEIGHTENED, CHECK-IN, ALERT at thresholds 0.38, 0.62 and 0.80. Only ALERT transmits. Biometric unlock is an automatic stand-down.',
])
P('Why not a rule table? Twelve binary signals is 2^12 = 4,096 cases to hand-author and maintain. '
  'A learned function over a feature vector is one function.', italic=True)

H('C · Zero-Knowledge Relay', 13, RUST, 6)
P('Identity is an ECDSA plus ECDH P-256 keypair generated on the device. Contacts pair by QR code. '
  'Alerts are sealed with ECDH to HKDF-SHA256 to AES-256-GCM, signed, and posted with a random '
  'per-alert recipient tag so two alerts from the same person cannot be linked.')
P('The reference relay implementation is under 60 lines of Python. It has no user table because '
  'there are no users. Seize the server and you obtain a queue of ciphertext and a rate-limit counter.')

H('D · Escalation ladder', 13, RGBColor(0xC0,0x8A,0x2A), 6)
table(['T+', 'Audience', 'Condition'],
 [['0 s', 'Device only', 'silent arming, breadcrumb buffer begins'],
  ['10 s', 'The user', 'silent vibrate; tap, biometric or voice cancels'],
  ['30 s', 'Nearby guardians', 'verified users within 400 m'],
  ['60 s', 'Trusted contacts', 'live location plus 10-minute breadcrumb'],
  ['90 s', '112 / police', 'requires three or more corroborating layers']],
 [0.7, 1.9, 4.0])
P('A false alarm is one friend receiving a notification. A missed alarm is a life. We therefore tune '
  'for sensitivity and make cancelling trivial.', italic=True, color=RUST)

H('E · City Data Layer', 13, RGBColor(0x4E,0x9E,0x7E), 6)
P('Aggregate and anonymous. Ranks interventions by risk reduction per rupee against a streetlight '
  'budget, using a counterfactual re-score of each segment as if it were lit. Buyers are municipal '
  'corporations, transport authorities, night-shift employers and campuses.')
P('The app helps one woman avoid one street tonight. The dashboard fixes the street so that nobody '
  'has to avoid it next year.', italic=True)

doc.add_page_break()
H('4 · The alert sequence', 16)
P('This is the only sequence in the entire product that causes a network request.')
image('docs/img/alert_sequence.png', 'Figure 2 — Steps 1 to 4 never leave the device. Step 5 is the first byte transmitted, and it is ciphertext.', 6.6)

H('5 · Code architecture', 16)
P('Two rules keep the codebase legible, and everything else follows from them.')
bullets([
 'core/ never imports React. It is pure, testable logic that would port to React Native or a Node worker unchanged.',
 'No screen imports an API implementation. Screens import the contract; the build decides whether that resolves to the in-browser adapter or an HTTP client.',
])
image('docs/img/code_map.png', 'Figure 3 — Module layout. services/api/contract.ts is the seam where a real backend plugs in.', 6.6)

doc.add_page_break()
H('6 · API contract', 16)
P('The entire backend surface is five endpoints, specified in docs/openapi.yaml. No endpoint accepts '
  'or returns anything that identifies a person. There is no /users, no /login and no /location.')
table(['Endpoint', 'Purpose', 'Contains identity?'],
 [['POST /risk/tiles', 'scored street segments for a bbox and hour', 'no'],
  ['POST /routes', 'Pareto-optimal routes over time vs risk', 'no'],
  ['POST /relay/alert', 'submit an opaque encrypted envelope', 'no — ciphertext only'],
  ['GET /city/works', 'interventions ranked by risk reduction per rupee', 'no — aggregate'],
  ['GET /places', 'open, lit, staffed places near a point', 'no']],
 [1.8, 3.3, 1.5])
P('The same contract is satisfied in the browser today, which is why the product deploys as a static '
  'bundle and keeps working in airplane mode. Switching to a real server is one environment variable.')
code('VITE_API_MODE=local   # default — computed in the browser, static deploy\n'
     'VITE_API_MODE=http    # talk to a FastAPI service implementing openapi.yaml')

H('7 · Deployment', 16)
image('docs/img/deployment.png', 'Figure 4 — Deployment topology. The dashed services are optional; the product works without them.', 6.6)
table(['Target', 'Command', 'Result'],
 [['Local development', 'npm run dev', 'http://localhost:5173'],
  ['Static build', 'npm run build', 'dist/ — deploy anywhere'],
  ['Netlify / Vercel', 'git push', 'netlify.toml / vercel.json included'],
  ['GitHub Pages', 'git push to main', '.github/workflows/deploy.yml'],
  ['Docker', 'npm run docker', 'nginx on :8080, optional relay on :8000']],
 [1.6, 2.0, 3.0])
P('Sensors require a secure context. The accelerometer, geolocation and microphone will not '
  'initialise over plain HTTP on a LAN address. Deploy to HTTPS or use a tunnel, then open on a phone.')
code('Permissions-Policy: accelerometer=(self), gyroscope=(self),\n'
     '                    geolocation=(self), microphone=(self)')

doc.add_page_break()
H('8 · Threat model', 16)
H('Mitigated', 12, TEAL, 4)
bullets([
 'Server compromise — there is nothing of value stored to steal.',
 'Network interception — payloads are end-to-end encrypted and signed.',
 'Alert spoofing — envelopes are ECDSA signed and rate limited per tag.',
 'Stalkerware repurposing — no silent monitoring mode, a persistent armed indicator, contacts removable without notification, and a duress PIN that shows a convincing all-clear while alerting.',
])
H('Not mitigated — stated openly', 12, RUST, 4)
bullets([
 'A compromised operating system defeats any user-space application.',
 'The relay still observes connection metadata such as IP address and timing.',
 'A phone that is already powered off cannot alert; we can only alert on the silence.',
 'Indoor GPS is unreliable, which degrades trajectory features.',
 'A familiar route with a known attacker produces no behavioural anomaly.',
])
P('We never claim to prevent crime. We claim to reduce time-to-detection from hours to roughly '
  '90 seconds, with zero user action.', italic=True, color=EBONY)

H('9 · What is real in the current build', 16)
table(['Component', 'Status'],
 [['Chennai street graph (2,970 segments, 213 POIs)', 'Real — OpenStreetMap, bundled'],
  ['Risk model r(segment, hour)', 'Real algorithm, provenance shown'],
  ['Pareto routing (Dijkstra, 6-alpha sweep)', 'Real'],
  ['Accelerometer feature extraction', 'Real — DeviceMotion on any phone browser'],
  ['GPS speed and vehicle onset', 'Real — watchPosition'],
  ['Scream detection', 'Real — Web Audio, on-device, nothing recorded'],
  ['Autoencoder anomaly model', 'Real — trained, weights committed'],
  ['Fusion and five-state machine', 'Real'],
  ['ECDH + AES-256-GCM sealing and signing', 'Real — Web Crypto'],
  ['Offline operation and SMS payload', 'Real — service worker, sms: intent'],
  ['VIIRS, Mapillary CNN, GNN, PU-learning', 'Roadmap'],
  ['Production relay, guardian network, BLE density', 'Roadmap']],
 [4.1, 2.5])

f = doc.add_paragraph(); f.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = f.add_run('Street data © OpenStreetMap contributors, ODbL · Built for the Crescent Project Expo')
r.font.size = Pt(8.5); r.font.color.rgb = INK3
f.paragraph_format.space_before = Pt(24)

doc.save('docs/NIZHAL_Architecture.docx')
import os; print('saved docs/NIZHAL_Architecture.docx', os.path.getsize('docs/NIZHAL_Architecture.docx')//1024, 'KB')
